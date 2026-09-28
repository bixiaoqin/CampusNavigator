# -*- coding: utf-8 -*-
"""
增强答辩PPT：添加算法/路网/数据存储说明页
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

PPT_PATH = r"C:\Users\lenovo\WorkBuddy\2026-06-17-20-24-23\CampusNavigator\docs\CampusNavigator_答辩.pptx"
DIAGRAM_DIR = r"C:\Users\lenovo\WorkBuddy\2026-06-17-20-24-23\CampusNavigator\docs\diagrams"

prs = Presentation(PPT_PATH)
print(f"原有幻灯片数: {len(prs.slides)}")
print(f"幻灯片尺寸: {prs.slide_width} x {prs.slide_height}")

# 使用空白布局
blank_layout = prs.slide_layouts[6]  # 空白
slide = prs.slides.add_slide(blank_layout)

# === 背景 ===
bg = slide.background
fill = bg.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0x1a, 0x1a, 0x2e)  # 深色背景

# === 标题 ===
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "核心算法 · 路网构建 · 数据存储"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = RGBColor(0x64, 0xB5, 0xF6)
p.alignment = PP_ALIGN.CENTER

# === 三列布局 ===
col_w = Inches(3.1)
col_y = Inches(1.3)
col_h = Inches(5.2)
gap = Inches(0.2)

# --- 第一列：算法 ---
col1_x = Inches(0.3)
# 列标题
txBox = slide.shapes.add_textbox(col1_x, col_y, col_w, Inches(0.5))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "① Dijkstra 最短路径算法"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(0x4C, 0xAF, 0x50)

# 列内容
txBox = slide.shapes.add_textbox(col1_x, col_y + Inches(0.5), col_w, Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True
items = [
    ("数据结构:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  • 邻接表 unordered_map<int, vector<Edge>>", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • 最小堆 priority_queue (懒惰删除)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • dist[] 距离表 + prev[] 前驱表", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("", False, None),
    ("算法流程:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  1. 初始化 dist[起点]=0, 其余=∞", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  2. 优先队列出队距离最小节点", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  3. 松弛所有邻居 (更新更短距离)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  4. 重复直到到达终点", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  5. 沿 prev[] 回溯最短路径", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("", False, None),
    ("复杂度: O((V+E) log V)", True, RGBColor(0xFF, 0xAB, 0x40)),
    ("V=66节点, E=80有向边", False, RGBColor(0xB0, 0xBE, 0xC5)),
    ("", False, None),
    ("创新: 天气感知Dijkstra", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  雨天→遮棚路权重×0.7", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  高温→树荫路权重递减", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  大风→临湖路权重×1.5", False, RGBColor(0xE0, 0xE0, 0xE0)),
]
for i, (text, bold, color) in enumerate(items):
    if i == 0:
        para = tf.paragraphs[0]
    else:
        para = tf.add_paragraph()
    para.text = text
    para.font.size = Pt(11)
    if bold:
        para.font.bold = True
    if color:
        para.font.color.rgb = color
    para.space_after = Pt(2)

# --- 第二列：路网 ---
col2_x = col1_x + col_w + gap
txBox = slide.shapes.add_textbox(col2_x, col_y, col_w, Inches(0.5))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "② 路网构建方式"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(0x42, 0xA5, 0xF5)

txBox = slide.shapes.add_textbox(col2_x, col_y + Inches(0.5), col_w, Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True
items2 = [
    ("底图来源:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  天津理工大学真实校园地图", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  tjut_map.jpg (1280×976px)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("", False, None),
    ("节点设计:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  • 45栋建筑 (id 0-44)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("    nodeKind=1, 显示标签", False, RGBColor(0xB0, 0xBE, 0xC5)),
    ("  • 21个路口节点 (id 100-120)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("    nodeKind=0, 仅参与寻路", False, RGBColor(0xB0, 0xBE, 0xC5)),
    ("  共66个节点", False, RGBColor(0xFF, 0xAB, 0x40)),
    ("", False, None),
    ("道路布局:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  • 大成路 x≈640 南北主干道", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • 知行道 y≈850 东西主干道", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • 横向道路 y=440, y=660", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • 40条无向道路 (双向加边)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("", False, None),
    ("边权重计算:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  weight = √(dx² + dy²)", False, RGBColor(0xCE, 0x93, 0xD8)),
    ("  (两端坐标欧氏距离)", False, RGBColor(0xB0, 0xBE, 0xC5)),
    ("", False, None),
    ("天气标签:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  shelter/tree/lake", False, RGBColor(0xE0, 0xE0, 0xE0)),
]
for i, (text, bold, color) in enumerate(items2):
    if i == 0:
        para = tf.paragraphs[0]
    else:
        para = tf.add_paragraph()
    para.text = text
    para.font.size = Pt(11)
    if bold:
        para.font.bold = True
    if color:
        para.font.color.rgb = color
    para.space_after = Pt(2)

# --- 第三列：数据存储 ---
col3_x = col2_x + col_w + gap
txBox = slide.shapes.add_textbox(col3_x, col_y, col_w, Inches(0.5))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "③ 数据存储方案"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(0xFF, 0x70, 0x43)

txBox = slide.shapes.add_textbox(col3_x, col_y + Inches(0.5), col_w, Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True
items3 = [
    ("双层架构:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  内存层: STL邻接表 (O(V+E))", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  持久层: SQLite 3 数据库", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("", False, None),
    ("SQLite表结构:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  buildings表 (9字段):", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("   id|name|x|y|type|info", False, RGBColor(0xCE, 0x93, 0xD8)),
    ("   |hours|floors|nodeKind", False, RGBColor(0xCE, 0x93, 0xD8)),
    ("  roads表 (3字段):", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("   from_id|to_id|weight", False, RGBColor(0xCE, 0x93, 0xD8)),
    ("", False, None),
    ("DatabaseManager:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  • 单例模式 (全局唯一实例)", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • QSqlQuery预处理语句", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • 防SQL注入", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  • 完整CRUD接口", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("", False, None),
    ("数据加载流程:", True, RGBColor(0xFF, 0xD5, 0x4F)),
    ("  启动→init()→建表", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  →空则seedData()", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  →loadAll→Graph邻接表", False, RGBColor(0xE0, 0xE0, 0xE0)),
    ("  →MapScene渲染", False, RGBColor(0xE0, 0xE0, 0xE0)),
]
for i, (text, bold, color) in enumerate(items3):
    if i == 0:
        para = tf.paragraphs[0]
    else:
        para = tf.add_paragraph()
    para.text = text
    para.font.size = Pt(11)
    if bold:
        para.font.bold = True
    if color:
        para.font.color.rgb = color
    para.space_after = Pt(2)

# === 底部路网图 ===
img_path = os.path.join(DIAGRAM_DIR, "road_network_diagram.png")
if os.path.exists(img_path):
    # 添加小图到底部
    pic = slide.shapes.add_picture(img_path, Inches(0.5), Inches(6.6), width=Inches(9))
    print("已添加路网图到PPT")

# 保存
output_path = PPT_PATH  # 覆盖原文件
prs.save(output_path)
print(f"PPT增强完成！新增算法/路网/数据存储页，总幻灯片数: {len(prs.slides)}")
